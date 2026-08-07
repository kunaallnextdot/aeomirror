"""Unit tests for bulk-scan URL parsing + validation (pure, DNS bypassed in tests):
normalize/dedupe, invalid rejection, over-limit truncation, SSRF blocking, CSV/XLSX
first-column extraction with header skipping, and the single-domain policy."""
import io

import openpyxl

import app.scanner.bulk as bulk_mod
from app.config import settings
from app.scanner.bulk import (
    build_url_list, extract_urls_from_file, is_single_registered_domain,
    parse_csv, parse_json, parse_text, parse_tsv, parse_xlsx,
)
from tests.authutil import auth_client


def test_normalize_and_dedupe():
    urls = ["example.com", "https://example.com/", "http://www.example.com", "example.com/a"]
    accepted, skipped = build_url_list(urls, 50)
    assert len(accepted) == 2                                    # example.com + /a
    assert [s["reason"] for s in skipped] == ["duplicate", "duplicate"]


def test_invalid_urls_skipped():
    accepted, skipped = build_url_list(["not a url", "", "https://ok.com/"], 50)
    assert accepted == ["https://ok.com/"]
    assert any(s["reason"] == "invalid" for s in skipped)


def test_stray_forbidden_char_is_invalid_not_ssrf():
    """A URL with a raw RFC-3986-forbidden char (a stray trailing '<') is rejected as
    `invalid` (via normalize_url → None), not mislabeled `ssrf_blocked`. Encoded %3C
    stays valid."""
    accepted, skipped = build_url_list(["https://a.com/x<", "https://ok.com/a%3Cb"], 50)
    assert accepted == ["https://ok.com/a%3Cb"]
    assert any(s["reason"] == "invalid" for s in skipped)
    assert not any(s["reason"] == "ssrf_blocked" for s in skipped)


def test_over_limit_truncates_to_skipped():
    urls = [f"https://ex{i}.com/" for i in range(60)]
    accepted, skipped = build_url_list(urls, 50)
    assert len(accepted) == 50
    assert sum(1 for s in skipped if s["reason"] == "over_limit") == 10


def test_ssrf_blocked_url_is_skipped(monkeypatch):
    # Re-enable the real SSRF check (tests bypass it by default) so a private/link-local
    # address is blocked. 169.254.169.254 is the classic cloud-metadata SSRF target.
    from app.core import ssrf
    monkeypatch.setattr(ssrf.settings, "allow_private_hosts", False)
    accepted, skipped = build_url_list(["http://169.254.169.254/latest/", "https://example.com/"], 50)
    assert "https://example.com/" in accepted
    assert any(s["reason"] == "ssrf_blocked" for s in skipped)


def test_xlsx_first_column_skips_header():
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.append(["URL", "Notes"])              # header row (first cell isn't a URL)
    ws.append(["https://a.com/", "home"])
    ws.append(["b.com/pricing", "extra"])
    buf = io.BytesIO()
    wb.save(buf)
    assert parse_xlsx(buf.getvalue()) == ["https://a.com/", "b.com/pricing"]


def test_xlsx_no_header_keeps_first_row():
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.append(["https://a.com/"])            # first cell IS a URL → not a header
    ws.append(["https://b.com/"])
    buf = io.BytesIO()
    wb.save(buf)
    assert parse_xlsx(buf.getvalue()) == ["https://a.com/", "https://b.com/"]


def test_csv_first_column_skips_header():
    data = b"URL,notes\nhttps://a.com/,home\nb.com/x,y\n"
    assert parse_csv(data) == ["https://a.com/", "b.com/x"]


def test_single_registered_domain_policy():
    assert is_single_registered_domain(["https://a.com/", "https://www.a.com/x", "http://a.com/y"])
    assert not is_single_registered_domain(["https://a.com/", "https://b.com/"])


# ------------------------------- text / json / tsv parsers (Steps 3) -------------------------------
def test_parse_text_lines_prose_and_comments():
    data = (b"https://a.com/\n"
            b"# a comment line\n"
            b"Visit https://b.com/pricing for details\n"
            b"c.com/x   # trailing note\n"
            b"\n")
    assert parse_text(data) == ["https://a.com/", "https://b.com/pricing", "c.com/x"]


def test_parse_text_single_comma_and_tab_line():
    assert parse_text(b"a.com, b.com, c.com\n") == ["a.com", "b.com", "c.com"]
    assert parse_text(b"a.com\tb.com\tc.com\n") == ["a.com", "b.com", "c.com"]


def test_parse_json_array_of_strings():
    assert parse_json(b'["https://a.com/", "b.com/x"]') == ["https://a.com/", "b.com/x"]


def test_parse_json_array_of_objects():
    data = b'[{"url": "https://a.com/"}, {"Link": "b.com/x"}, {"note": "skip"}]'
    assert parse_json(data) == ["https://a.com/", "b.com/x"]


def test_parse_json_object_with_urls_key():
    assert parse_json(b'{"urls": ["https://a.com/", "b.com/x"]}') == ["https://a.com/", "b.com/x"]


def test_parse_jsonl_one_object_per_line():
    data = b'{"url": "https://a.com/"}\n{"url": "https://b.com/"}\n'
    assert parse_json(data) == ["https://a.com/", "https://b.com/"]


def test_parse_json_malformed_falls_back_to_text():
    # Not JSON at all → never raises; treated as one URL per line.
    assert parse_json(b"https://a.com/\nhttps://b.com/\n") == ["https://a.com/", "https://b.com/"]


def test_parse_tsv_first_column_skips_header():
    data = b"URL\tnotes\nhttps://a.com/\thome\nb.com/x\ty\n"
    assert parse_tsv(data) == ["https://a.com/", "b.com/x"]


# ------------------------------- encoding + content sniffing (Steps 2–5) -------------------------------
def test_csv_utf16_is_decoded():
    text = "URL,notes\nhttps://a.com/,home\nb.com/x,y\n"
    assert parse_csv(text.encode("utf-16")) == ["https://a.com/", "b.com/x"]   # BOM included


def test_csv_cp1252_is_decoded():
    # A non-ASCII cp1252 byte (é = \xe9) must not derail decoding of the ASCII URLs.
    text = "URL,notes\nhttps://a.com/,caf\xe9\nb.com/x,y\n"
    assert parse_csv(text.encode("cp1252")) == ["https://a.com/", "b.com/x"]


def test_mislabelled_xlsx_named_csv_parsed_by_content():
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.append(["https://a.com/"])
    ws.append(["https://b.com/"])
    buf = io.BytesIO(); wb.save(buf)
    # Real XLSX bytes, but a .csv filename → content sniffing (PK magic) must win.
    assert extract_urls_from_file("list.csv", buf.getvalue()) == ["https://a.com/", "https://b.com/"]


def test_xlsx_urls_in_third_column_are_found():
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.append(["Name", "Region", "URL"])          # header; column A has no URLs
    ws.append(["Acme", "EU", "https://a.com/"])
    ws.append(["Beta", "US", "b.com/pricing"])
    buf = io.BytesIO(); wb.save(buf)
    assert parse_xlsx(buf.getvalue()) == ["https://a.com/", "b.com/pricing"]


def test_legacy_ods_rejected_with_supported_hint():
    import zipfile
    buf = io.BytesIO()
    with zipfile.ZipFile(buf, "w") as z:
        z.writestr("mimetype", "application/vnd.oasis.opendocument.spreadsheet")
    try:
        extract_urls_from_file("sheet.ods", buf.getvalue())
        assert False, "expected ValueError for a legacy .ods upload"
    except ValueError as e:
        assert ".csv" in str(e) and ".xlsx" in str(e)


def test_legacy_xls_rejected_with_supported_hint():
    data = b"\xD0\xCF\x11\xE0\xA1\xB1\x1a\xe1" + b"\x00" * 40      # OLE2 magic (legacy .xls)
    try:
        extract_urls_from_file("book.xls", data)
        assert False, "expected ValueError for a legacy .xls upload"
    except ValueError as e:
        assert "re-save" in str(e).lower()


# ------------------------------- upload size cap (Step 1) -------------------------------
def test_oversized_upload_is_413_and_never_parsed(monkeypatch):
    """An upload over bulk_upload_max_bytes is rejected with 413 BEFORE the parser
    runs (so a zip-bomb .xlsx can't decompress and OOM the worker)."""
    monkeypatch.setattr(settings, "bulk_upload_max_bytes", 100)

    def _must_not_run(*a, **k):
        raise AssertionError("parser reached despite oversized upload")
    monkeypatch.setattr(bulk_mod, "extract_urls_from_file", _must_not_run)

    client, _ = auth_client()
    big = b"https://example.com/\n" * 50               # ~1 KB, well over the 100-byte cap
    r = client.post("/v1/scan/bulk", files={"file": ("urls.csv", big, "text/csv")})
    assert r.status_code == 413
    assert "too large" in r.json()["detail"].lower()
